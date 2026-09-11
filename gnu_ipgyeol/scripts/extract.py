# -*- coding: utf-8 -*-
"""경상국립대 수시 입학결과(2020~2026학년도)에서
농학과 / 원예(원예학과·원예과학부) / 식물의학과의
학생부교과(일반)·학생부종합(일반) 결과를 뽑아 CSV로 만든다.

원자료: gnu_ipgyeol/source/ (경상국립대 입학안내 공개 파일 원본)
출력:   gnu_ipgyeol/gnu_입결_2020_2026.csv
"""
import csv
import os
import re

import openpyxl
import pymupdf

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
SRC = os.path.join(BASE, "source")
OUT = os.path.join(BASE, "gnu_입결_2020_2026.csv")

TARGETS = ("농학과", "원예학과", "원예과학부", "농업식물과학과", "식물의학과")

COLUMNS = [
    "학년도", "전형유형", "전형명", "모집단위",
    "모집인원", "지원인원", "경쟁률", "최종등록인원", "충원최종예비순위",
    "학생부등급_평균", "상위누적_평균", "학생부등급_80컷", "상위누적_80컷",
    "환산점수_평균", "환산점수_80컷", "비고",
]

# 9등급제 표준 누적비율. 등급 사이는 직선 보간한다.
CUM = {1: 4, 2: 11, 3: 23, 4: 40, 5: 60, 6: 77, 7: 89, 8: 96, 9: 100}


def cum_pct(grade):
    """학생부등급을 상위 누적 백분위(%) 근사값으로 바꾼다."""
    if not grade:
        return ""
    g = float(grade)
    lo = max(1, min(8, int(g)))
    return f"{CUM[lo] + (g - lo) * (CUM[lo + 1] - CUM[lo]):.1f}%"


def num(v):
    """셀 값을 문자열로 정리한다(공백·None 제거, 소수는 둘째 자리까지)."""
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{round(v, 2):g}"
    s = str(v).strip()
    return "" if s in ("-", "None") else s


def rate(v):
    s = num(v)
    return re.sub(r"\s+", "", s)


def from_xlsx(year, gyogwa_sheet, jonghap_sheet, has_jeonhyeong_col):
    """2020~2022학년도: 엑셀 원자료."""
    wb = openpyxl.load_workbook(os.path.join(SRC, f"{year}_수시입학결과.xlsx"))
    rows = []

    ws = wb[gyogwa_sheet]
    for r in ws.iter_rows(values_only=True):
        r = list(r)
        off = 1 if has_jeonhyeong_col else 0
        unit = num(r[1 + off]) if len(r) > 1 + off else ""
        if unit not in TARGETS:
            continue
        if has_jeonhyeong_col and num(r[0]) != "교과(일반)":
            continue  # 지역인재(교과) 등은 제외
        rows.append([
            year, "학생부교과", "일반", unit,
            num(r[2 + off]), num(r[3 + off]), rate(r[4 + off]),
            num(r[8 + off]), num(r[9 + off]),
            num(r[11 + off]), cum_pct(num(r[11 + off])),
            num(r[12 + off]), cum_pct(num(r[12 + off])),
            num(r[13 + off]), num(r[14 + off]), "",
        ])

    ws = wb[jonghap_sheet]
    for r in ws.iter_rows(values_only=True):
        r = list(r)
        unit = num(r[2]) if len(r) > 2 else ""
        if unit not in TARGETS:
            continue
        rows.append([
            year, "학생부종합", "일반", unit,
            num(r[3]), num(r[4]), rate(r[5]), num(r[6]), num(r[7]),
            num(r[9]), cum_pct(num(r[9])), "", "",
            num(r[10]), num(r[11]), "",
        ])
    return rows


def from_pdf(year, filename):
    """2023~2026학년도: PDF 원자료(표 추출)."""
    doc = pymupdf.open(os.path.join(SRC, filename))
    rows = []
    for page in doc:
        for table in page.find_tables().tables:
            for raw in table.extract():
                c = [(x or "").strip() for x in raw]
                if not c:
                    continue
                # 2026학년도에만 '정원구분' 열이 하나 더 있다.
                if len(c) > 1 and c[1] in ("정원내", "정원외"):
                    c = c[:1] + c[2:]
                mode = c[0].replace("수시(", "").replace(")", "")
                if mode not in ("학생부교과", "학생부종합") or len(c) < 5:
                    continue
                if c[1] != "일반" or c[3] not in TARGETS:
                    continue
                unit = c[3]
                if mode == "학생부교과":  # 수능최저 관련 3개 열이 중간에 있다
                    rows.append([
                        year, mode, "일반", unit,
                        c[4], c[5], c[6], c[10], c[11],
                        c[12], cum_pct(c[12]), c[13], cum_pct(c[13]),
                        c[14], c[15], "",
                    ])
                else:
                    rows.append([
                        year, mode, "일반", unit,
                        c[4], c[5], c[6], c[7], c[8],
                        c[9], cum_pct(c[9]), "", "", "", "", "",
                    ])
    return rows


DEPT_GROUPS = [
    ("농학과", ("농학과", "농업식물과학과")),
    ("원예과학부(구 원예학과)", ("원예과학부", "원예학과")),
    ("식물의학과", ("식물의학과",)),
]

HEAD_GYOGWA = ["학년도", "모집단위", "모집인원", "지원인원", "경쟁률", "최종등록",
               "충원최종예비", "등급 평균", "상위 누적", "등급 80컷", "상위 누적",
               "환산 평균", "환산 80컷"]
HEAD_JONGHAP = ["학년도", "모집단위", "모집인원", "지원인원", "경쟁률", "최종등록",
                "충원최종예비", "등급 평균", "상위 누적", "환산 평균", "환산 80컷"]


def md_table(header, body):
    out = ["| " + " | ".join(header) + " |",
           "|" + "|".join(["---"] * len(header)) + "|"]
    for r in body:
        out.append("| " + " | ".join(x if x else "-" for x in r) + " |")
    return "\n".join(out)


def write_markdown(rows):
    doc = [
        "# 경상국립대학교 수시 입학결과 - 농학과 / 원예과학부 / 식물의학과",
        "",
        "학생부교과(일반전형)과 학생부종합(일반전형) 기준, 2020~2026학년도 7개년.",
        "경상국립대학교 입학안내 홈페이지가 공개한 학년도별 수시모집 입학결과 원본을",
        "`source/` 에 그대로 두고, `scripts/extract.py` 로 뽑아낸 값이다.",
        "",
        "## 읽기 전에",
        "",
        "- 최종등록자(입학생) 기준이며 학생부등급은 대학 자체 환산 방식으로 산출한 값이다.",
        "- 80컷은 등록자를 성적순으로 줄 세웠을 때 80% 지점의 값이다.",
        "- 환산점수 산출식은 학년도·전형별로 달라 연도 간 직접 비교에 적합하지 않다. 추세는 학생부등급으로 본다.",
        "- 학생부종합 전형은 2023학년도부터 등급 평균만 공개하고 등급 80컷·환산점수를 공개하지 않는다.",
        "- '상위 누적'은 9등급제 표준 누적비율(1등급 4%, 2등급 11%, 3등급 23%, 4등급 40%, 5등급 60% …)을",
        "  등급 사이에서 직선 보간해 등급을 백분위로 바꾼 근사치다. 대학이 발표한 값이 아니다.",
        "- 2020학년도에는 농학과와 원예학과가 '농업식물과학과'로 통합 모집됐고, 2021학년도에 분리됐다.",
        "- 원예학과는 2022학년도부터 원예과학부로 개편됐고, 2025·2026학년도에는 학생부종합 일반전형을 선발하지 않았다.",
        "",
    ]
    for title, units in DEPT_GROUPS:
        doc.append(f"## {title}")
        doc.append("")
        for mode, head, idx in (
            ("학생부교과", HEAD_GYOGWA, [0, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14]),
            ("학생부종합", HEAD_JONGHAP, [0, 3, 4, 5, 6, 7, 8, 9, 10, 13, 14]),
        ):
            body = []
            for r in sorted(rows, key=lambda r: r[0]):
                if r[1] != mode or r[3] not in units:
                    continue
                line = [str(r[i]) for i in idx]
                if r[15].startswith("해당 학년도 미선발"):
                    line = [line[0], line[1]] + ["미선발"] * (len(line) - 2)
                body.append(line)
            doc.append(f"### {mode}(일반전형)")
            doc.append("")
            doc.append(md_table(head, body))
            doc.append("")
    doc += [
        "## 원자료",
        "",
        "| 학년도 | 파일 | 출처 |",
        "|---|---|---|",
        "| 2020 | source/2020_수시입학결과.xlsx | https://www.gnu.ac.kr/new/cm/cntnts/cntntsView.do?mi=5003&cntntsId=2866 |",
        "| 2021 | source/2021_수시입학결과.xlsx | https://www.gnu.ac.kr/new/cm/cntnts/cntntsView.do?mi=4996&cntntsId=2862 |",
        "| 2022 | source/2022_수시입학결과.xlsx | https://www.gnu.ac.kr/new/cm/cntnts/cntntsView.do?mi=10920&cntntsId=5299 |",
        "| 2023 | source/2023_수시입학결과.pdf | https://www.gnu.ac.kr/new/cm/cntnts/cntntsView.do?mi=13223&cntntsId=6234 |",
        "| 2024 | source/2024_수시입학결과.pdf | https://new.gnu.ac.kr/new/cm/cntnts/cntntsView.do?mi=17447&cntntsId=7958 |",
        "| 2025 | source/2025_수시입학결과.pdf | https://new.gnu.ac.kr/new/cm/cntnts/cntntsView.do?mi=19078&cntntsId=8709 |",
        "| 2026 | source/2026_수시입학결과.pdf | https://new.gnu.ac.kr/new/cm/cntnts/cntntsView.do?mi=48029&cntntsId=21841 |",
        "",
        "## 다시 만들기",
        "",
        "```bash",
        "pip install pymupdf openpyxl",
        "python3 gnu_ipgyeol/scripts/extract.py",
        "```",
        "",
    ]
    with open(os.path.join(BASE, "README.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(doc))


def main():
    rows = []
    rows += from_xlsx(2020, "학생부교과 및 실기", "학생부종합(일반전형)", False)
    rows += from_xlsx(2021, "학생부교과 및 실기", "학생부종합(일반전형)", False)
    rows += from_xlsx(2022, "학생부 교과 및 실기", "학생부종합(일반)", True)
    rows += from_pdf(2023, "2023_수시입학결과.pdf")
    rows += from_pdf(2024, "2024_수시입학결과.pdf")
    rows += from_pdf(2025, "2025_수시입학결과.pdf")
    rows += from_pdf(2026, "2026_수시입학결과.pdf")

    # 원예과학부는 2025·2026학년도 학생부종합 일반전형을 선발하지 않았다.
    for year in (2025, 2026):
        rows.append([
            year, "학생부종합", "일반", "원예과학부",
            "", "", "", "", "", "", "", "", "", "", "",
            "해당 학년도 미선발(모집요강·결과자료에 모집단위 없음)",
        ])

    for r in rows:
        if r[3] == "농업식물과학과":
            r[15] = "2021학년도에 농학과·원예학과로 분리되기 전 통합 모집단위"
        elif r[3] == "원예학과" and r[0] <= 2021:
            r[15] = "2022학년도부터 원예과학부로 개편"

    order = {"농학과": 0, "농업식물과학과": 0, "원예학과": 1, "원예과학부": 1, "식물의학과": 2}
    rows.sort(key=lambda r: (order[r[3]], r[1], r[0]))

    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        w.writerows(rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "입결"
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, col in enumerate(COLUMNS, start=1):
        width = 46 if col == "비고" else max(10, len(col) + 4)
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    wb.save(os.path.join(BASE, "gnu_입결_2020_2026.xlsx"))

    write_markdown(rows)
    print(f"{len(rows)}행 저장: {OUT}")


if __name__ == "__main__":
    main()
